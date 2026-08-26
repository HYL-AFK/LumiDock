from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "LumiDock-已用引脚表-无颜色.docx"
FONT = "Microsoft YaHei"
BLACK = "000000"
GRAY = "000000"


def set_run(run, size=10.5, bold=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=90, start=130, bottom=90, end=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    for tag, value in (("w:tblW", sum(widths)), ("w:tblInd", 120)):
        item = tblPr.first_child_found_in(tag)
        if item is None:
            item = OxmlElement(tag)
            tblPr.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
    layout = tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, rows):
    widths = (1500, 3000, 4860)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    headers = ("引脚", "对应功能", "位置与接法")
    for cell, text in zip(table.rows[0].cells, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 10, True)
    for pin, function, location in rows:
        cells = table.add_row().cells
        values = (pin, function, location)
        for index, (cell, text) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            set_run(p.add_run(text), 9.4)
    return table


def heading(doc, text):
    p = doc.add_paragraph()
    p.style = "Heading 1"
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), 14, True)


def body(doc, text, size=10.5, color=BLACK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.16
    set_run(p.add_run(text), size, False, color)


def make_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16
    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = FONT
    heading_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading_style.font.color.rgb = RGBColor.from_string(BLACK)
    heading_style.font.size = Pt(14)
    heading_style.font.bold = True
    heading_style.paragraph_format.space_before = Pt(10)
    heading_style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    set_run(footer.add_run("LumiDock | 当前已用引脚表"), 8.5, False, GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run(title.add_run("LumiDock 当前已用引脚表"), 20, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    set_run(subtitle.add_run("ESP32-S3 Super Mini | 黑白无颜色版本 | 2026-07-31"), 9.5, False, GRAY)

    body(doc, "表格只列当前项目已经分配功能的引脚，不包含 GPIO0、GPIO3、GPIO19/20、GPIO43、GPIO45、GPIO46、GPIO48 等保留或调试引脚。第三列同时说明开发板焊盘位置和外部接线对象。")

    rows = [
        ("3V3", "ESP32 逻辑电源", "开发板外侧电源焊盘；接 3.3 V 逻辑外设，不能接 5 V 信号"),
        ("GND", "系统公共地", "开发板外侧 GND 焊盘；与屏幕、WS2812、音频、RS485、充电模块共地"),
        ("GPIO1", "WS2812 #1 数据", "开发板正面下排外侧焊盘；外接第一路 WS2812，5 V 灯珠需电平转换"),
        ("GPIO2", "WS2812 #2 数据", "开发板正面下排外侧焊盘；外接第二路 WS2812，5 V 灯珠需电平转换"),
        ("GPIO4", "I2S BCLK", "开发板正面下排外侧焊盘；外接 I2S 麦克风和 I2S 功放共用时钟"),
        ("GPIO5", "I2S LRCLK / WS", "开发板正面下排外侧焊盘；外接 I2S 麦克风和功放共用帧时钟"),
        ("GPIO6", "I2S MIC_DATA", "开发板正面下排外侧焊盘；连接数字 I2S 麦克风数据输出"),
        ("GPIO7", "I2S AMP_DIN", "开发板正面下排外侧焊盘；连接 I2S 功放数据输入"),
        ("GPIO8", "ST7789 RES", "开发板正面上排外侧焊盘；外接 ST7789 屏幕复位"),
        ("GPIO9", "ST7789 DC", "开发板正面上排外侧焊盘；外接 ST7789 屏幕数据/命令选择"),
        ("GPIO10", "ST7789 CS", "开发板正面上排外侧焊盘；外接 ST7789 屏幕片选"),
        ("GPIO11", "ST7789 SDA / MOSI", "开发板正面上排外侧焊盘；外接 ST7789 SPI 数据"),
        ("GPIO12", "ST7789 SCL / SCLK", "开发板正面上排外侧焊盘；外接 ST7789 SPI 时钟"),
        ("GPIO13", "ST7789 BLK", "开发板正面上排外侧焊盘；外接屏幕背光，可使用 PWM"),
        ("GPIO14", "单按键", "开发板背面左侧内侧焊盘；外接按键，按键另一端接 GND，内部上拉"),
        ("GPIO15", "I2C SDA", "开发板背面左侧内侧焊盘；连接 SW3518S Pin 20 SDA，也可共享其他 I2C 设备"),
        ("GPIO16", "I2C SCK / SCL", "开发板背面左侧内侧焊盘；连接 SW3518S Pin 21 SCK，I2C 地址 0x3C"),
        ("GPIO17", "MOS1 控制", "开发板背面左侧内侧焊盘；焊线接 MOS 驱动/电子开关 1 的控制端"),
        ("GPIO18", "MOS2 控制", "开发板背面左侧内侧焊盘；焊线接 MOS 驱动/电子开关 2 的控制端"),
        ("GPIO21", "单线 UART", "开发板背面左侧内侧焊盘；焊线接外部单线 UART，电平和收发器按协议确定"),
        ("GPIO37", "RS485 TX", "开发板背面左侧内侧焊盘；焊线接自动收发 RS485 模块 TX/DI 方向"),
        ("GPIO38", "RS485 RX", "开发板背面左侧内侧焊盘；焊线接自动收发 RS485 模块 RX/RO 方向"),
        ("GPIO39", "雷达 UART_RX", "开发板背面右侧内侧焊盘；ESP32 接收雷达 TX 数据"),
        ("GPIO40", "雷达 UART_TX", "开发板背面右侧内侧焊盘；ESP32 发送到雷达 RX 数据"),
        ("GPIO41", "雷达 OUT", "开发板背面右侧内侧焊盘；外接雷达数字输出，5 V 输出必须先降到 3.3 V"),
    ]
    add_table(doc, rows)

    heading(doc, "接线备注")
    body(doc, "1. GPIO15/GPIO16 同一条 I2C 总线当前用于读取 SW3518S。SW3518S 的 SDA/SCK 逻辑电平和上拉电压需要确认，ESP32 GPIO 不能直接承受 5 V。")
    body(doc, "2. GPIO17/GPIO18 只输出 MOS 控制信号，不直接承载外部负载电流。MOS 类型、负载电压、电流和续流保护需要按实际负载设计。")
    body(doc, "3. GPIO21 按单线 UART 使用，不按标准 1-Wire 使用；不要默认增加 4.7 kOhm 1-Wire 上拉。")
    body(doc, "4. WS2812 采用外接供电。5 V 灯珠的数据线建议使用 74AHCT1G125/74HCT125 电平转换，并在数据线上串联约 330 Ohm 电阻。")
    body(doc, "5. 当前表格中的“外侧焊盘”表示开发板边缘可接线焊盘；“内侧焊盘”表示没有常规排针插槽，需焊线、转接板或自制排针引出。", 9.5, GRAY)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_doc()
