from pathlib import Path
from tempfile import gettempdir

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "LumiDock-业务流程与实现架构.docx"
ASSET_DIR = Path(gettempdir()) / "lumidock_business_flow_assets"
FONT = "Microsoft YaHei"
FONT_FILE = Path("C:/Windows/Fonts/msyh.ttc")
BLACK = "000000"
GRAY = "555555"


def pil_font(size, bold=False):
    index = 1 if bold else 0
    return ImageFont.truetype(str(FONT_FILE), size=size, index=index)


def text_center(draw, box, text, font, spacing=8):
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        ((left + right - width) / 2, (top + bottom - height) / 2 - 2),
        text,
        font=font,
        fill="black",
        spacing=spacing,
        align="center",
    )


def box(draw, rect, text, size=30, bold=False):
    draw.rectangle(rect, outline="black", width=3, fill="white")
    text_center(draw, rect, text, pil_font(size, bold))


def decision(draw, rect, text, size=28):
    left, top, right, bottom = rect
    cx, cy = (left + right) / 2, (top + bottom) / 2
    draw.polygon([(cx, top), (right, cy), (cx, bottom), (left, cy)], outline="black", fill="white", width=3)
    text_center(draw, (left + 70, top + 28, right - 70, bottom - 28), text, pil_font(size, True), spacing=6)


def arrow(draw, start, end, width=3):
    draw.line([start, end], fill="black", width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        head = [(x2, y2), (x2 - 18 * direction, y2 - 9), (x2 - 18 * direction, y2 + 9)]
    else:
        direction = 1 if y2 > y1 else -1
        head = [(x2, y2), (x2 - 9, y2 - 18 * direction), (x2 + 9, y2 - 18 * direction)]
    draw.polygon(head, fill="black")


def elbow_arrow(draw, points):
    for start, end in zip(points, points[1:]):
        draw.line([start, end], fill="black", width=3)
    start, end = points[-2], points[-1]
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        head = [(x2, y2), (x2 - 18 * direction, y2 - 9), (x2 - 18 * direction, y2 + 9)]
    else:
        direction = 1 if y2 > y1 else -1
        head = [(x2, y2), (x2 - 9, y2 - 18 * direction), (x2 + 9, y2 - 18 * direction)]
    draw.polygon(head, fill="black")


def label(draw, xy, text, size=25):
    draw.text(xy, text, font=pil_font(size, True), fill="black")


def canvas(width, height, title):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 34), title, font=pil_font(38, True), fill="black")
    draw.line([(60, 98), (width - 60, 98)], fill="black", width=3)
    return image, draw


def make_main_flow():
    image, draw = canvas(2100, 3220, "LumiDock 主业务流程")
    center_x = 1050
    vertical = [
        ((710, 140, 1390, 270), "上电"),
        ((710, 330, 1390, 490), "GPIO 安全初始化\nMOS 关闭、功放静音、灯效低功耗"),
        ((710, 550, 1390, 710), "初始化驱动\n屏幕、灯带、I2S、RS485、雷达、一线通"),
    ]
    for rect, text in vertical:
        box(draw, rect, text, 31, rect == vertical[0][0])
    arrow(draw, (center_x, 270), (center_x, 330))
    arrow(draw, (center_x, 490), (center_x, 550))

    test = (770, 790, 1330, 980)
    decision(draw, test, "硬件自检\n通过？")
    arrow(draw, (center_x, 710), (center_x, 790))
    fault = (1490, 810, 2010, 960)
    box(draw, fault, "故障模式\n关闭危险输出、记录故障、重试", 27)
    elbow_arrow(draw, [(1330, 885), (1410, 885), (1410, 885), (1490, 885)])
    label(draw, (1370, 825), "否", 24)

    cfg = (710, 1050, 1390, 1200)
    box(draw, cfg, "读取配置\nWi-Fi、云端、灯效、阈值、设备参数", 29)
    arrow(draw, (center_x, 980), (center_x, 1050))
    wifi = (770, 1280, 1330, 1470)
    decision(draw, wifi, "Wi-Fi\n已连接？")
    arrow(draw, (center_x, 1200), (center_x, 1280))
    offline = (1490, 1300, 2010, 1450)
    box(draw, offline, "离线待机\n本地功能继续运行\n后台指数退避重连", 27)
    elbow_arrow(draw, [(1330, 1375), (1410, 1375), (1410, 1375), (1490, 1375)])
    label(draw, (1370, 1315), "否", 24)

    cloud = (770, 1540, 1330, 1730)
    decision(draw, cloud, "云端鉴权\n连接成功？")
    arrow(draw, (center_x, 1470), (center_x, 1540))
    elbow_arrow(draw, [(1330, 1635), (1410, 1635), (1410, 1375), (1490, 1375)])
    label(draw, (1370, 1575), "否", 24)
    ready = (710, 1810, 1390, 1950)
    box(draw, ready, "待机 IDLE\n低亮或息屏，等待外部事件", 31, True)
    arrow(draw, (center_x, 1730), (center_x, 1810))

    label(draw, (70, 2025), "事件输入", 32)
    inputs = [
        ((70, 2100, 405, 2230), "雷达\n靠近/离开"),
        ((440, 2100, 775, 2230), "按键\n短按/双击/长按"),
        ((810, 2100, 1145, 2230), "RS485 / 一线通\n数据更新、超时"),
        ((1180, 2100, 1515, 2230), "云端命令\n控制、配置、OTA"),
        ((1550, 2100, 2030, 2230), "定时器 / 硬件异常\n超时、重连、故障"),
    ]
    for rect, text in inputs:
        box(draw, rect, text, 24)
        elbow_arrow(draw, [((rect[0] + rect[2]) / 2, 2100), ((rect[0] + rect[2]) / 2, 2020), (center_x, 2020), (center_x, 1950)])

    queue = (620, 2320, 1480, 2460)
    box(draw, queue, "事件中心 Event Queue\n统一来源、时间戳、优先级", 30, True)
    arrow(draw, (center_x, 2230), (center_x, 2320))
    state = (620, 2540, 1480, 2710)
    box(draw, state, "状态机与规则仲裁\n故障保护 > 本地按键 > 自动规则 > 云端/语音 > 展示动画", 26, True)
    arrow(draw, (center_x, 2460), (center_x, 2540))
    valid = (800, 2780, 1300, 2960)
    decision(draw, valid, "命令合法且\n允许执行？", 28)
    arrow(draw, (center_x, 2710), (center_x, 2780))
    reject = (90, 2800, 630, 2940)
    box(draw, reject, "拒绝执行\n提示原因，必要时上报", 27)
    elbow_arrow(draw, [(800, 2870), (710, 2870), (710, 2870), (630, 2870)])
    label(draw, (700, 2805), "否", 24)
    command = (1440, 2800, 2010, 2940)
    box(draw, command, "生成统一设备命令\nCommand Queue", 28, True)
    elbow_arrow(draw, [(1300, 2870), (1370, 2870), (1370, 2870), (1440, 2870)])
    label(draw, (1340, 2805), "是", 24)
    execbox = (480, 3035, 1620, 3170)
    box(draw, execbox, "执行服务：屏幕 / WS2812 / 音频 / MOS / RS485 / 数据缓存 / NVS\n执行结果 -> 本地反馈 + 云端状态/事件/命令回执 -> 回到待机", 27, True)
    elbow_arrow(draw, [(1725, 2940), (1725, 3000), (1050, 3000), (1050, 3035)])
    elbow_arrow(draw, [(1050, 3035), (1050, 2990), (250, 2990), (250, 2950)])
    image.save(ASSET_DIR / "main_flow.png")
    return ASSET_DIR / "main_flow.png"


def make_voice_flow():
    image, draw = canvas(2100, 1720, "云端 ASR 语音子流程")
    stages = [
        ((710, 140, 1390, 280), "按键长按触发语音"),
        ((710, 350, 1390, 510), "状态检查\n网络在线、未故障、未升级、未占用"),
    ]
    for rect, text in stages:
        box(draw, rect, text, 30, rect == stages[0][0])
    arrow(draw, (1050, 280), (1050, 350))
    check = (770, 580, 1330, 770)
    decision(draw, check, "允许进入\n语音会话？", 29)
    arrow(draw, (1050, 510), (1050, 580))
    deny = (1490, 600, 2010, 750)
    box(draw, deny, "提示不可用\n回到当前状态", 28)
    elbow_arrow(draw, [(1330, 675), (1410, 675), (1410, 675), (1490, 675)])
    label(draw, (1370, 615), "否", 24)
    audio = (710, 850, 1390, 1020)
    box(draw, audio, "播放提示音，启动 I2S 麦克风\n采集 16 kHz / 单声道 / 16 bit PCM", 29)
    arrow(draw, (1050, 770), (1050, 850))
    stop = (770, 1090, 1330, 1280)
    decision(draw, stop, "说话结束或\n达到录音上限？", 29)
    arrow(draw, (1050, 1020), (1050, 1090))
    elbow_arrow(draw, [(770, 1185), (580, 1185), (580, 935), (710, 935)])
    label(draw, (610, 1125), "否", 24)
    cloud = (710, 1360, 1390, 1510)
    box(draw, cloud, "上传音频到云端 ASR\n得到文本 -> 指令解析 -> 安全校验 -> 执行 -> 屏幕/语音反馈", 28, True)
    arrow(draw, (1050, 1280), (1050, 1360))
    image.save(ASSET_DIR / "voice_flow.png")
    return ASSET_DIR / "voice_flow.png"


def make_fault_flow():
    image, draw = canvas(2100, 1750, "异常、离线与恢复流程")
    start = (710, 140, 1390, 280)
    box(draw, start, "检测到异常或通信超时", 32, True)
    kinds = [
        ((60, 460, 510, 630), "Wi-Fi / 云端断开\n进入离线模式\n后台重连"),
        ((570, 460, 1030, 630), "RS485 / 一线通超时\n标记外设离线\n周期轮询恢复"),
        ((1080, 460, 1540, 630), "屏幕 / 灯 / I2S 异常\n关闭对应服务\n有限次数重初始化"),
        ((1590, 460, 2050, 630), "MOS / 供电风险\n立即关闭 MOS\n停止高功耗输出"),
    ]
    for rect, text in kinds:
        box(draw, rect, text, 25)
        elbow_arrow(draw, [(1050, 280), (1050, 370), ((rect[0] + rect[2]) / 2, 370), ((rect[0] + rect[2]) / 2, 460)])
    recover = (560, 800, 1540, 960)
    box(draw, recover, "统一故障处理\n屏幕告警、写入故障码、在线时上报、避免重复告警", 30, True)
    for rect, _ in kinds:
        elbow_arrow(draw, [((rect[0] + rect[2]) / 2, 630), ((rect[0] + rect[2]) / 2, 720), (1050, 720), (1050, 800)])
    nextstate = (770, 1100, 1330, 1290)
    decision(draw, nextstate, "是否已恢复\n且自检通过？", 29)
    arrow(draw, (1050, 960), (1050, 1100))
    retry = (80, 1120, 600, 1270)
    box(draw, retry, "保持降级状态\n延时后再次尝试恢复", 28)
    elbow_arrow(draw, [(770, 1195), (690, 1195), (690, 1195), (600, 1195)])
    label(draw, (650, 1135), "否", 24)
    returnidle = (710, 1400, 1390, 1540)
    box(draw, returnidle, "恢复服务，补发关键状态\n回到在线或离线待机", 30, True)
    arrow(draw, (1050, 1290), (1050, 1400))
    image.save(ASSET_DIR / "fault_flow.png")
    return ASSET_DIR / "fault_flow.png"


def set_run(run, size=10.5, bold=False, color=BLACK, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag, value in (("w:tblW", str(sum(widths))), ("w:tblInd", "120")):
        item = tblPr.first_child_found_in(tag)
        if item is None:
            item = OxmlElement(tag)
            tblPr.append(item)
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
    layout = tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_col, width in zip(tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 9.3, True)
    for values in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            set_run(p.add_run(text), 9.0)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), 15 if level == 1 else 12.5, True)
    return p


def body(doc, text, prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    if prefix and text.startswith(prefix):
        set_run(p.add_run(prefix), 10.5, True)
        set_run(p.add_run(text[len(prefix):]), 10.5)
    else:
        set_run(p.add_run(text), 10.5)
    return p


def make_doc():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    main_flow = make_main_flow()
    voice_flow = make_voice_flow()
    fault_flow = make_fault_flow()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.74)
    section.right_margin = Inches(0.74)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12.5, 10, 5)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    set_run(footer.add_run("LumiDock | 业务流程与实现架构"), 8.5, False, GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run(title.add_run("LumiDock 业务流程与实现架构"), 20, True)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    set_run(sub.add_run("基于 ESP32-S3 Super Mini 的本地感知、语音、屏幕、灯效、外设与云端协同方案 | 2026-07-30"), 9.5, False, GRAY)

    heading(doc, "1. 项目定位")
    body(doc, "LumiDock 是一个带屏幕、双路外接 WS2812、I2S 麦克风与功放、雷达、RS485、一线通、双 MOS 输出和云端连接能力的智能终端。软件的核心是把不同来源的输入统一为事件，再经过状态机和规则仲裁执行设备动作。")
    body(doc, "设计原则：", "设计原则：")
    body(doc, "1) 任何输入都不能直接操作 GPIO；必须先进入事件中心。")
    body(doc, "2) 云端断开不影响本地感知、按键、屏幕、灯效和外设控制。")
    body(doc, "3) MOS 和外部设备控制必须经过安全校验；复位、异常和供电风险时默认关闭。")
    body(doc, "4) 一线通当前按数据/状态接口抽象；其具体帧格式和控制方向需在设备协议确认后补充。")

    heading(doc, "2. 软件分层")
    add_table(doc, ("层级", "职责", "包含模块"), [
        ("硬件驱动层", "只负责外设初始化、读写和中断。", "ST7789、WS2812、I2S、RS485、雷达、一线通、按键、MOS、I2C"),
        ("设备服务层", "将驱动包装成可被业务层调用的服务。", "UI、灯效、音频、感知、RS485、数据、配置、网络、云端、OTA"),
        ("事件与命令层", "统一来源、优先级和执行结果。", "Event Queue、Command Queue、状态缓存、故障记录"),
        ("业务状态机", "决定当前设备状态和各输入的可执行性。", "启动、自检、在线待机、离线待机、语音会话、故障、升级"),
        ("云端协议层", "处理鉴权、命令、遥测、回执和配置。", "MQTT/HTTPS、设备状态、事件、命令确认、OTA"),
    ], (1700, 3600, 4060))

    heading(doc, "3. 主业务流程图")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Main flow is intentionally narrower so the complete diagram fits on one page.
    p.add_run().add_picture(str(main_flow), width=Inches(5.45))
    body(doc, "主流程说明：", "主流程说明：")
    body(doc, "上电后先把所有输出置为安全状态，再初始化并自检硬件。网络或云端失败时进入离线待机，而不是停止运行。运行中的所有输入统一进入 Event Queue，由状态机检查优先级、当前状态和安全条件，再产生 Command Queue 中的设备命令。")

    heading(doc, "4. 输入事件、决策与输出")
    add_table(doc, ("输入来源", "事件示例", "状态机处理", "典型输出"), [
        ("雷达", "有人靠近、离开超时", "唤醒、维持活跃、回到待机", "亮屏、切换灯效、更新在线状态"),
        ("单按键", "短按、双击、长按", "本地快捷动作或启动语音会话", "切换页面、执行预设、开始/结束录音"),
        ("云端", "控制、配置更新、OTA", "验证设备身份、命令编号和状态许可", "执行控制、保存配置、回传回执"),
        ("RS485 / 一线通", "数据更新、通信超时", "更新状态或标记外设离线", "刷新页面、遥测上报、周期重试"),
        ("定时器", "离开超时、周期上报、重连", "执行自动化规则", "息屏、关灯、重连、上报"),
        ("硬件故障", "通信失败、供电风险、任务异常", "优先进入保护和降级", "关闭危险输出、告警、重初始化"),
    ], (1400, 1900, 3000, 3060))

    heading(doc, "5. 云端 ASR 语音子流程")
    body(doc, "第一版建议使用按键长按启动语音，不依赖本地唤醒词。这样能先验证云端 ASR、I2S 音频链路、网络稳定性和命令安全策略；离线 ESP-SR 可在确认 Flash/PSRAM 后再作为第二阶段能力。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(voice_flow), width=Inches(6.85))
    add_table(doc, ("阶段", "要求"), [
        ("会话许可", "网络在线、未处于故障或 OTA、未被其他录音任务占用。"),
        ("音频采集", "I2S 麦克风采集 16 kHz、单声道、16 bit PCM，按固定帧长缓存。"),
        ("云端识别", "通过 HTTPS 或 WebSocket 上传音频；需要设置录音上限、网络超时和取消机制。"),
        ("命令解析", "先将文本映射到受控命令表，未识别或未支持的文本只反馈，不执行外部动作。"),
        ("高风险动作", "MOS、RS485 控制等命令必须再过状态机与安全规则，不能仅凭 ASR 文本直接执行。"),
    ], (1850, 7510))

    heading(doc, "6. 异常、离线与恢复流程")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(fault_flow), width=Inches(6.85))
    body(doc, "离线模式：", "离线模式：")
    body(doc, "Wi-Fi 或云端中断后，本地业务继续运行；网络服务采用指数退避策略重连。重连成功后，上报当前状态和关键事件，不盲目重放可能失效的旧控制命令。")
    body(doc, "故障模式：", "故障模式：")
    body(doc, "对于 MOS、供电和可能造成外部设备风险的异常，先执行保护动作，再展示和上报故障。对于单个外围设备异常，优先关闭对应服务并降级运行，避免一个外设失败导致主机反复重启。")

    heading(doc, "7. 优先级与互斥规则")
    add_table(doc, ("优先级", "来源", "规则"), [
        ("1", "硬件保护、故障状态", "可立即关闭 MOS、功放和高功耗灯效；覆盖任何普通控制命令。"),
        ("2", "本地按键", "本地用户明确操作优先于展示和自动灯效；长按可以中断当前语音会话。"),
        ("3", "自动化规则", "雷达离开超时、休眠、周期维护等按配置执行，但不能突破安全状态。"),
        ("4", "云端与语音命令", "经身份、参数、状态和安全校验后执行；冲突时保留最新且未过期的合法命令。"),
        ("5", "展示任务", "页面轮播、动态灯效和提示动画可被任意更高优先级事件打断。"),
    ], (1100, 2300, 5960))

    heading(doc, "8. 推荐实现阶段")
    stages = [
        "阶段 1：完成 GPIO 和外设驱动的独立测试，包括 ST7789、双 WS2812、按键、MOS、雷达、RS485、I2S。",
        "阶段 2：实现 Event Queue、Command Queue、设备状态缓存和启动/自检/待机状态机。",
        "阶段 3：打通本地闭环：雷达唤醒、按键操作、屏幕状态、灯效、MOS 安全控制。",
        "阶段 4：接入 RS485 与一线通协议，完成数据缓存、通信超时和本地展示。",
        "阶段 5：实现 Wi-Fi 配网、云端协议、命令回执、离线重连和 NVS 配置持久化。",
        "阶段 6：加入按键触发的云端 ASR，再根据实际内存规格评估离线 ESP-SR。",
        "阶段 7：完成 OTA、看门狗、故障日志、掉电恢复、产测和长期稳定性测试。",
    ]
    for item in stages:
        body(doc, item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_doc()
