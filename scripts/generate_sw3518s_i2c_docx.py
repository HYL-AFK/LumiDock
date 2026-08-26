from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SW3518S_I2C通信与充电状态读取资料.docx"
FONT = "Microsoft YaHei"
MONO = "Consolas"
BLACK = "000000"
GRAY = "555555"


def set_run(run, size=10.5, bold=False, color=BLACK, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    for tag, value in (("w:tblW", sum(widths)), ("w:tblInd", 120)):
        item = tblPr.first_child_found_in(tag)
        if item is None:
            item = OxmlElement(tag)
            tblPr.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
    layout = tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths, font_size=9.1):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), font_size, True)
    for values in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(text), font_size)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), 15 if level == 1 else 12.5, True)
    return p


def body(doc, text, prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    if prefix and text.startswith(prefix):
        set_run(p.add_run(prefix), 10.5, True)
        set_run(p.add_run(text[len(prefix):]), 10.5)
    else:
        set_run(p.add_run(text), 10.5)
    return p


def code_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_run(p.add_run(text), 9.2, name=MONO)
    return p


def make_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12.5, 10, 5)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    set_run(footer.add_run("LumiDock | SW3518S I2C 通信资料"), 8.5, color=GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run(title.add_run("SW3518S I2C 通信与充电状态读取资料"), 20, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_run(subtitle.add_run("适用于 LumiDock ESP32-S3 子项目 | 依据 SW3518S 原厂数据手册与寄存器手册整理 | 2026-07-31"), 9.5, color=GRAY)

    heading(doc, "1. 结论")
    body(doc, "SW3518S 与 ESP32-S3 之间的主机通信接口确定为 I2C。芯片 20 脚为 SDA 数据线，21 脚为 SCK 时钟线；原厂数据手册明确说明 SCK 是 I2C clock。接口支持 100 kHz 和 400 kHz。")
    body(doc, "USB PD、PPS、QC、AFC、FCP、SCP、VOOC 等是 SW3518S 对外的快充协商协议，不是 ESP32 和 SW3518S 之间的通信协议。ESP32 只读取 SW3518S 已经协商出的协议和电气状态。")

    heading(doc, "2. 芯片与接口信息")
    add_table(doc, ("项目", "确认内容", "LumiDock 处理"), [
        ("芯片", "SW3518S，双口快充电源管理芯片", "作为被读取的 I2C 从机"),
        ("SDA", "芯片 Pin 20，I2C 数据", "连接 ESP32 GPIO15"),
        ("SCK", "芯片 Pin 21，I2C 时钟", "连接 ESP32 GPIO16；功能等同标准 I2C SCL"),
        ("速率", "100 kHz / 400 kHz", "第一版使用 100 kHz，稳定后再评估 400 kHz"),
        ("从机地址", "原厂写为 0x3C", "ESP-IDF 使用 7 位地址 0x3C，不要写成 0x78"),
    ], (1800, 3850, 3310))
    body(doc, "线连接参考：")
    code_line(doc, "ESP32 GPIO15  ---- SDA ---- SW3518S Pin 20")
    code_line(doc, "ESP32 GPIO16  ---- SCK/SCL - SW3518S Pin 21")
    code_line(doc, "ESP32 GND     ------------ SW3518S GND")
    body(doc, "SDA/SCK 的上拉电压和模块是否已经内置上拉，需要实测或结合 Demo 原理图确认。ESP32 不能直接承受 5 V I2C 电平；若模块上拉到 5 V，需要双向电平转换。")

    heading(doc, "3. I2C 读时序")
    body(doc, "原厂资料给出的读操作是先写入寄存器地址，再重复起始读取数据。0x3C 是 7 位从机地址；在物理字节层，写地址为 0x78，读地址为 0x79，但使用 ESP-IDF I2C API 时传入 0x3C。")
    code_line(doc, "START -> 0x78(W) -> ACK -> register -> ACK")
    code_line(doc, "RESTART -> 0x79(R) -> ACK -> data -> NACK -> STOP")
    body(doc, "第一阶段只做只读。不要写入配置寄存器、PD Hard Reset 寄存器或 PDO 配置寄存器。")

    heading(doc, "4. 版本识别与协议寄存器")
    add_table(doc, ("寄存器", "字段", "含义"), [
        ("0x01", "Bit 2:0", "IC Version；0x05 对应旧版寄存器手册，0x07 对应 RG004_2_v2.1"),
        ("0x06", "Bit 5:4", "PD 版本：1=PD 2.0，2=PD 3.0"),
        ("0x06", "Bit 3:0", "快充协议编码，见下一表"),
        ("0x07", "Bit 2", "Buck 状态：0=关闭，1=开启"),
        ("0x07", "Bit 1", "Type-A 端口状态"),
        ("0x07", "Bit 0", "Type-C 端口状态"),
        ("0x08", "Bit 7:4", "外部设备连接状态，适用于 A+C 端口模式"),
    ], (1600, 1900, 5460))
    add_table(doc, ("0x06[3:0]", "协议"), [
        ("0x1", "QC2.0"), ("0x2", "QC3.0"), ("0x3", "FCP"), ("0x4", "SCP"),
        ("0x5", "PD Fixed"), ("0x6", "PD PPS"), ("0x7", "PE1.1"), ("0x8", "PE2.0"),
        ("0x9", "VOOC"), ("0xA", "SFCP"), ("0xB", "AFC"),
    ], (1800, 7160), 9.0)

    heading(doc, "5. 电压、电流和温度读取")
    add_table(doc, ("寄存器", "分辨率", "含义"), [
        ("0x30", "160 mV/bit", "输入电压，快速读取"),
        ("0x31", "96 mV/bit", "输出电压，快速读取"),
        ("0x33", "40 mA/bit", "Type-C 输出电流，快速读取"),
        ("0x34", "40 mA/bit", "Type-A 输出电流，快速读取"),
        ("0x37", "8 mV/bit", "NTC 电压数据"),
    ], (1800, 1900, 5460))
    body(doc, "芯片还提供 12 位 ADC。通过 0x3A 选择 ADC 类型，再读取 0x3B 高 8 位和 0x3C 低 4 位。12 位 ADC 分辨率为：输入电压 10 mV/bit、输出电压 6 mV/bit、Type-C/Type-A 电流 2.5 mA/bit。第一版界面可先使用快速读取寄存器，后续再增加高精度显示。")

    heading(doc, "6. ESP32 软件读取流程")
    code_line(doc, "启动 -> 初始化 I2C(100 kHz) -> 读取 0x01 芯片版本")
    code_line(doc, "      -> 读取 0x06 协议 -> 读取 0x07 状态 -> 读取 0x08 设备连接")
    code_line(doc, "      -> 读取 0x30/0x31/0x33/0x34 -> 换算 V/A/W -> 更新 ChargingState")
    code_line(doc, "      -> 屏幕显示 -> 云端上报 -> 500 ms~1 s 后再次读取")
    body(doc, "建议建立 ChargingState 结构：protocol、pd_version、input_voltage_mv、output_voltage_mv、type_a_current_ma、type_c_current_ma、power_mw、port_status、data_valid、last_update_time。")
    body(doc, "读取失败时显示“通信异常”或“数据无效”，不要无提示地继续显示过期数据。")

    heading(doc, "7. 暂不修改的寄存器")
    add_table(doc, ("寄存器", "用途", "当前策略"), [
        ("0x12", "开启 0xA0~0xBF 配置写入权限", "第一阶段不写"),
        ("0x70", "PD 通用请求，可发送 Hard Reset", "禁止写入"),
        ("0xA0~0xBF", "功率、PDO、协议开关和端口配置", "只读研究，暂不改配置"),
    ], (1600, 3500, 4060))
    body(doc, "写配置前必须确认 SW3518S 芯片版本、当前产品功率设计和厂家参数，否则可能改变 PD/PPS 能力、最大电压、电流限制或端口行为。")

    heading(doc, "8. 资料来源与版本匹配")
    body(doc, "本资料依据厂家提供的 SW3518S Datasheet_Release_DS031_v1.2.pdf、SW3518S_Register_List_Release_RG004_2_v2.1.pdf 和 SW3518S 原理图设计指南整理。寄存器手册 v2.1 说明其对应 IC version 7；实际固件应先读 0x01，再选择对应寄存器定义。")
    body(doc, "SW3518S 设计指南还提示：如果使用外部 MCU，应从 VDRV 取电且外部器件电流尽量不超过 20 mA。ESP32-S3 的供电应使用 LumiDock 自己的 3.3 V 电源，不要用 VDRV 直接供电；I2C 只共享必要的信号和地。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_doc()
